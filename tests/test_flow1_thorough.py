"""
test_flow1_thorough.py (SIH PS 26117)
=====================================
Comprehensive, thorough test suite for P9 Flow 1 (Scanned Document -> OCR -> RAG -> DocGen Approval Note).

Test Scenarios Covered:
  Scenario 1: Standard Single-Page Scanned PDF Approval Note Generation (test-1.pdf)
  Scenario 2: Procurement Line-Item Extraction & Approval Note (test-2.pdf)
  Scenario 3: Custom Pre-indexed RAG Template Ingestion & CAPEX Approval Synthesis (test-3.pdf)
  Scenario 4: Programmatic Output Deliverable (.docx) Structure & Integrity Inspection
"""

import sys
import os
import uuid
import docx
from pathlib import Path
from typing import List, Dict, Any

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT_DIR)

from orchestrator import build_orchestrator_graph, DB_FILENAME, WorkbenchState
from router.schemas import FileMetadata
try:
    from scripts.inspect_run import inspect_thread
except ImportError:
    from scripts.inspect_run import inspect_thread
from tools.rag_kb import rag_kb, RagKbInput


def run_scenario_1() -> Dict[str, Any]:
    """Scenario 1: Standard Single-Page Scanned PDF Approval Note Generation (test-1.pdf)."""
    print("\n" + "=" * 75)
    print("RUNNING SCENARIO 1: Standard Single-Page PDF Approval Note (test-1.pdf)")
    print("=" * 75)

    thread_id = str(uuid.uuid4())
    pdf_path = str(Path("ignore-files/ocr-test/test-1.pdf").resolve())

    prompt = (
        f"I have a scanned report at {pdf_path}. Extract its contents, then generate a formal "
        "approval note in Word format based on our standard approval note template, incorporating "
        "the key details from the report. Save it as an approval note document."
    )

    initial_state = WorkbenchState(
        prompt=prompt,
        file_metadata=[
            FileMetadata(
                filename=Path(pdf_path).name,
                extension=Path(pdf_path).suffix,
                filepath=pdf_path,
                path=pdf_path,
            )
        ],
    )

    config = {"configurable": {"thread_id": thread_id}}
    builder = build_orchestrator_graph()

    with SqliteSaver.from_conn_string(DB_FILENAME) as checkpointer:
        compiled_graph = builder.compile(checkpointer=checkpointer)
        result = compiled_graph.invoke(initial_state, config=config)

    inspect_thread(thread_id)

    tool_results = result.get("tool_results", []) or []
    called_tools = [
        getattr(tr, "metadata", {}).get("tool_name") or "unknown" for tr in tool_results
    ]

    return {
        "scenario": "Scenario 1 (test-1.pdf)",
        "thread_id": thread_id,
        "called_tools": called_tools,
        "response": result.get("response"),
        "result": result,
    }


def run_scenario_2() -> Dict[str, Any]:
    """Scenario 2: Procurement Line-Item Extraction & Approval Note (test-2.pdf)."""
    print("\n" + "=" * 75)
    print("RUNNING SCENARIO 2: Procurement Line-Item Extraction (test-2.pdf)")
    print("=" * 75)

    thread_id = str(uuid.uuid4())
    pdf_path = str(Path("ignore-files/ocr-test/test-2.pdf").resolve())

    prompt = (
        f"Extract all line items, quantities, and pricing from the scanned invoice/report at {pdf_path}. "
        "Search our knowledge base for procurement approval thresholds, and generate a formal "
        "procurement approval note in Word format (.docx) with a table of items and total cost."
    )

    initial_state = WorkbenchState(
        prompt=prompt,
        file_metadata=[
            FileMetadata(
                filename=Path(pdf_path).name,
                extension=Path(pdf_path).suffix,
                filepath=pdf_path,
                path=pdf_path,
            )
        ],
    )

    config = {"configurable": {"thread_id": thread_id}}
    builder = build_orchestrator_graph()

    with SqliteSaver.from_conn_string(DB_FILENAME) as checkpointer:
        compiled_graph = builder.compile(checkpointer=checkpointer)
        result = compiled_graph.invoke(initial_state, config=config)

    inspect_thread(thread_id)

    tool_results = result.get("tool_results", []) or []
    called_tools = [
        getattr(tr, "metadata", {}).get("tool_name") or "unknown" for tr in tool_results
    ]

    return {
        "scenario": "Scenario 2 (test-2.pdf Procurement)",
        "thread_id": thread_id,
        "called_tools": called_tools,
        "response": result.get("response"),
        "result": result,
    }


def run_scenario_3() -> Dict[str, Any]:
    """Scenario 3: Custom Pre-indexed RAG Template Ingestion & CAPEX Approval (test-3.pdf)."""
    print("\n" + "=" * 75)
    print("RUNNING SCENARIO 3: Pre-indexed RAG Template Ingestion & CAPEX Approval (test-3.pdf)")
    print("=" * 75)

    # Pre-ingest a corporate CAPEX approval guideline document into ChromaDB
    template_doc_path = Path("reference_files/capex_template_guide.txt").resolve()
    template_doc_path.parent.mkdir(parents=True, exist_ok=True)
    template_content = (
        "# Corporate CAPEX Approval Note Standard Operating Procedure\n\n"
        "Section 1: Mandatory Approval Note Headings\n"
        "1. Executive Summary: High-level justification and total expenditure.\n"
        "2. Equipment & Technical Specifications: Detailed breakdown of items and quantities.\n"
        "3. Financial Impact & Budget Head: Verification of cost allocation.\n"
        "4. Risk Assessment & Safety Compliance: Environmental and operational risks.\n"
        "5. Recommendation & Sign-off: Final recommendation for executive approval.\n\n"
        "Section 2: Approval Thresholds\n"
        "Expenditures above $50,000 require Board approval. Expenditures under $50,000 require Plant Manager approval.\n"
    )
    with open(template_doc_path, "w", encoding="utf-8") as f:
        f.write(template_content)

    print(f"[RAG Pre-ingest] Ingesting template guide at '{template_doc_path}'...")
    ingest_res = rag_kb(
        RagKbInput(
            action="ingest",
            file_path=str(template_doc_path),
            session_id="workbench_session",
        )
    )
    print(f"[RAG Pre-ingest] Ingestion result status: {ingest_res.status.value}")

    thread_id = str(uuid.uuid4())
    pdf_path = str(Path("ignore-files/ocr-test/test-3.pdf").resolve())

    prompt = (
        f"I have a scanned equipment report at {pdf_path}. First extract its text using ocr_vlm. "
        "Then query our RAG knowledge base for CAPEX approval note headings and guidelines. "
        "Finally, generate a formal CAPEX approval note in Word format (.docx) following the corporate "
        "CAPEX template structure."
    )

    initial_state = WorkbenchState(
        prompt=prompt,
        file_metadata=[
            FileMetadata(
                filename=Path(pdf_path).name,
                extension=Path(pdf_path).suffix,
                filepath=pdf_path,
                path=pdf_path,
            )
        ],
    )

    config = {"configurable": {"thread_id": thread_id}}
    builder = build_orchestrator_graph()

    with SqliteSaver.from_conn_string(DB_FILENAME) as checkpointer:
        compiled_graph = builder.compile(checkpointer=checkpointer)
        result = compiled_graph.invoke(initial_state, config=config)

    inspect_thread(thread_id)

    tool_results = result.get("tool_results", []) or []
    called_tools = [
        getattr(tr, "metadata", {}).get("tool_name") or "unknown" for tr in tool_results
    ]

    return {
        "scenario": "Scenario 3 (test-3.pdf CAPEX Template)",
        "thread_id": thread_id,
        "called_tools": called_tools,
        "response": result.get("response"),
        "result": result,
    }


def run_scenario_4_doc_audit() -> Dict[str, Any]:
    """Scenario 4: Programmatic Output Deliverable (.docx) Structure & Integrity Inspection."""
    print("\n" + "=" * 75)
    print("RUNNING SCENARIO 4: Programmatic Output (.docx) Structure & Integrity Audit")
    print("=" * 75)

    workspace_dir = Path("workspace").resolve()
    docx_files = list(workspace_dir.rglob("*.docx"))

    audit_results = []
    for doc_path in docx_files:
        try:
            doc = docx.Document(str(doc_path))
            p_count = len(doc.paragraphs)
            t_count = len(doc.tables)
            s_count = len(doc.sections)
            headings = [p.text for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]
            file_size_kb = round(doc_path.stat().st_size / 1024, 2)

            is_valid = file_size_kb > 0 and (p_count > 0 or t_count > 0)
            audit_results.append({
                "file_name": doc_path.name,
                "relative_path": str(doc_path.relative_to(workspace_dir)),
                "size_kb": file_size_kb,
                "paragraphs": p_count,
                "tables": t_count,
                "headings_count": len(headings),
                "headings": headings[:3],
                "valid": is_valid,
            })
            print(f"[Doc Audit] File: '{doc_path.name}' | Size: {file_size_kb} KB | Paragraphs: {p_count} | Tables: {t_count} | Valid: {is_valid}")
        except Exception as e:
            audit_results.append({
                "file_name": doc_path.name,
                "relative_path": str(doc_path),
                "error": str(e),
                "valid": False,
            })
            print(f"[Doc Audit ERROR] File: '{doc_path.name}' | Error: {e}")

    return {
        "total_files_audited": len(docx_files),
        "audit_results": audit_results,
        "all_valid": all(r.get("valid", False) for r in audit_results) if audit_results else False,
    }


def main():
    print("=" * 75)
    print("STARTING THOROUGH P9 FLOW 1 SUITE (4 TEST SCENARIOS)")
    print("=" * 75)

    s1 = run_scenario_1()
    s2 = run_scenario_2()
    s3 = run_scenario_3()
    s4 = run_scenario_4_doc_audit()

    print("\n" + "=" * 75)
    print("FLOW 1 THOROUGH SUITE - FINAL AUDIT SUMMARY MATRIX")
    print("=" * 75)

    scenarios = [s1, s2, s3]
    all_passed = True

    for idx, sc in enumerate(scenarios, 1):
        name = sc["scenario"]
        tools = sc["called_tools"]
        res = sc["result"]
        reached_end = res.get("response") is not None
        has_ocr = "ocr_vlm" in tools
        has_rag = "rag_kb" in tools
        has_doc = "doc_gen" in tools
        scenario_pass = reached_end and has_ocr and has_doc

        if not scenario_pass:
            all_passed = False

        print(f"Scenario {idx}: {name}")
        print(f"  - Tools Called Sequence : {tools}")
        print(f"  - OCR VLM Fired         : {'PASS' if has_ocr else 'FAIL'}")
        print(f"  - RAG KB Fired          : {'PASS' if has_rag else 'FAIL'}")
        print(f"  - DocGen Fired          : {'PASS' if has_doc else 'FAIL'}")
        print(f"  - Reached END Node      : {'PASS' if reached_end else 'FAIL'}")
        print(f"  - Scenario Verdict      : {'PASS' if scenario_pass else 'FAIL'}\n")

    print(f"Scenario 4: Document Integrity Audit")
    print(f"  - Total .docx Audited   : {s4['total_files_audited']}")
    print(f"  - All Deliverables Valid: {'PASS' if s4['all_valid'] else 'FAIL'}\n")

    print("=" * 75)
    if all_passed and s4["all_valid"]:
        print(">>> FLOW 1 THOROUGH TEST SUITE RESULT: PASSED 100% (ALL SCENARIOS SUCCESS)")
    else:
        print(">>> FLOW 1 THOROUGH TEST SUITE RESULT: PARTIAL / ISSUES DETECTED")
    print("=" * 75)


if __name__ == "__main__":
    main()

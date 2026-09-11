"""
Tool: doc_gen (SIH PS 26117)
============================
Output deliverable generator layer for Word (.docx), PowerPoint (.pptx), and Excel (.xlsx).
Uses format-agnostic content primitives (TextBlock, TableBlock, ImageBlock, ChartBlock)
composed into format-specific specs (DocxSpec, PptxSpec, XlsxSpec).
"""

import re
from pathlib import Path
from typing import Union, Optional, Any
from pydantic import BaseModel, Field, model_validator
import docx
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from tool_interface import ToolInput, ToolResult, ToolStatus, audited_tool, validate_workspace_path


# --- Shared Content Primitives ---

class TextBlock(BaseModel):
    heading: Optional[str] = None
    body: str


class TableBlock(BaseModel):
    headers: list[str]
    rows: list[list[str]]


class ImageBlock(BaseModel):
    path: str
    caption: Optional[str] = None


class ChartBlock(BaseModel):
    chart_type: str = "bar"  # "bar" | "line" | "pie"
    data: dict[str, list[Any]] = Field(default_factory=dict)
    title: Optional[str] = None


# --- Format-Specific Top-Level Specs ---

class DocxSpec(BaseModel):
    title: str
    sections: list[Union[TextBlock, TableBlock, ImageBlock]]


class SlideSpec(BaseModel):
    title: str
    blocks: list[Union[TextBlock, TableBlock, ImageBlock, ChartBlock]]


class PptxSpec(BaseModel):
    slides: list[SlideSpec]


class SheetSpec(BaseModel):
    name: str
    data: list[list[Any]]
    chart: Optional[ChartBlock] = None


class XlsxSpec(BaseModel):
    sheets: list[SheetSpec]


# --- Tool Input & Result Models ---

class DocGenInput(ToolInput):
    """
    Enterprise Document Deliverable Generator.
    Produces formatted Word documents (.docx), PowerPoint slide decks (.pptx), and Excel spreadsheets (.xlsx).
    Call this tool whenever the user asks to generate, format, create, or export an approval note, technical memo, operational report, presentation, or spreadsheet.
    """
    format: str = Field(default="docx", description="Document format: 'docx', 'pptx', or 'xlsx'")
    spec: dict = Field(default_factory=dict, description="Document specification dictionary")
    output_path: str = Field(default="approval_note.docx", description="Output document filename or relative path")
    session_id: str = Field(default="workbench_session", description="Session workspace identifier")

    @model_validator(mode="before")
    @classmethod
    def normalize_fields(cls, data: Any) -> Any:
        if isinstance(data, dict):
            # 1. Normalize format / output_format / word format alias
            if "format" not in data and "output_format" in data:
                data["format"] = data.pop("output_format")
            if data.get("format") in ("word", "doc", "document"):
                data["format"] = "docx"
            if data.get("format") in ("ppt", "pptx", "powerpoint", "slides", "presentation"):
                data["format"] = "pptx"
            if data.get("format") in ("excel", "sheet", "spreadsheet", "xls"):
                data["format"] = "xlsx"

            # 2. Normalize spec from docx_spec / pptx_spec / xlsx_spec / sections
            if "spec" not in data or not data["spec"]:
                for alt_key in ("docx_spec", "pptx_spec", "xlsx_spec", "sections", "slides"):
                    if alt_key in data:
                        alt_val = data.get(alt_key)
                        if alt_key == "sections" and isinstance(alt_val, list):
                            data["spec"] = {"title": "Generated Document", "sections": alt_val}
                        elif alt_key == "slides" and isinstance(alt_val, list):
                            data["spec"] = {"slides": alt_val}
                        else:
                            data["spec"] = alt_val
                        break

            # 3. Default output_path if missing or empty
            if not data.get("output_path"):
                spec_dict = data.get("spec", {})
                title = spec_dict.get("title") if isinstance(spec_dict, dict) else None
                fmt = data.get("format", "docx")
                if title and isinstance(title, str) and title.strip():
                    clean_title = re.sub(r'[^\w\s-]', '', title).strip()
                    data["output_path"] = f"{clean_title}.{fmt}"
                else:
                    data["output_path"] = f"approval_note.{fmt}"
        return data


class DocGenResult(ToolResult):
    output_path: Optional[str] = None
    format: Optional[str] = None


# --- Renderers ---

def _render_docx(spec: DocxSpec, output_file: Path, session_id: str):
    doc = docx.Document()
    doc.add_heading(spec.title, level=0)

    for block in spec.sections:
        if isinstance(block, TextBlock):
            if block.heading:
                doc.add_heading(block.heading, level=1)
            doc.add_paragraph(block.body)

        elif isinstance(block, TableBlock):
            table = doc.add_table(rows=1, cols=len(block.headers))
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            for idx, header in enumerate(block.headers):
                hdr_cells[idx].text = str(header)

            for row_data in block.rows:
                row_cells = table.add_row().cells
                for idx, cell_val in enumerate(row_data):
                    if idx < len(row_cells):
                        row_cells[idx].text = str(cell_val)

        elif isinstance(block, ImageBlock):
            try:
                img_path = validate_workspace_path(block.path, session_id, create_parents=False)
                if img_path.exists():
                    doc.add_picture(str(img_path), width=Inches(5))
                    if block.caption:
                        p = doc.add_paragraph(block.caption)
                        p.runs[0].font.italic = True
            except Exception as e:
                doc.add_paragraph(f"[Image Error: {e}]")

    doc.save(str(output_file))


def _render_pptx(spec: PptxSpec, output_file: Path, session_id: str):
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # Blank slide

    for slide_spec in spec.slides:
        slide = prs.slides.add_slide(blank_layout)
        
        # Add slide title
        txBox = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(0.5), PptxInches(8.4), PptxInches(1))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = slide_spec.title
        p.font.size = PptxPt(28)
        p.font.bold = True

        top_offset = 1.6
        for block in slide_spec.blocks:
            if isinstance(block, TextBlock):
                tb = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(top_offset), PptxInches(8.4), PptxInches(1.5))
                tf2 = tb.text_frame
                tf2.word_wrap = True
                if block.heading:
                    p_head = tf2.paragraphs[0]
                    p_head.text = block.heading
                    p_head.font.size = PptxPt(20)
                    p_head.font.bold = True
                    p_body = tf2.add_paragraph()
                else:
                    p_body = tf2.paragraphs[0]
                p_body.text = block.body
                p_body.font.size = PptxPt(16)
                top_offset += 1.8

            elif isinstance(block, TableBlock):
                rows = len(block.rows) + 1
                cols = len(block.headers)
                table_shape = slide.shapes.add_table(rows, cols, PptxInches(0.8), PptxInches(top_offset), PptxInches(8.4), PptxInches(1.8))
                table = table_shape.table

                for c_idx, header in enumerate(block.headers):
                    table.cell(0, c_idx).text = str(header)

                for r_idx, row_data in enumerate(block.rows):
                    for c_idx, val in enumerate(row_data):
                        if c_idx < cols:
                            table.cell(r_idx + 1, c_idx).text = str(val)

                top_offset += 2.2

            elif isinstance(block, ImageBlock):
                try:
                    img_path = validate_workspace_path(block.path, session_id, create_parents=False)
                    if img_path.exists():
                        slide.shapes.add_picture(str(img_path), PptxInches(0.8), PptxInches(top_offset), width=PptxInches(4.5))
                        top_offset += 3.0
                except Exception:
                    top_offset += 0.5

    prs.save(str(output_file))


def _render_xlsx(spec: XlsxSpec, output_file: Path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    for sheet_spec in spec.sheets:
        ws = wb.create_sheet(title=sheet_spec.name)
        for row in sheet_spec.data:
            ws.append(row)

        if sheet_spec.chart:
            c = sheet_spec.chart
            if c.chart_type == "line":
                chart_obj = LineChart()
            elif c.chart_type == "pie":
                chart_obj = PieChart()
            else:
                chart_obj = BarChart()

            chart_obj.title = c.title or "Data Chart"
            max_row = len(sheet_spec.data)
            max_col = len(sheet_spec.data[0]) if sheet_spec.data else 1
            data_ref = Reference(ws, min_col=2, min_row=1, max_col=max_col, max_row=max_row)
            cats_ref = Reference(ws, min_col=1, min_row=2, max_row=max_row)
            chart_obj.add_data(data_ref, titles_from_data=True)
            chart_obj.set_categories(cats_ref)
            ws.add_chart(chart_obj, "E2")

    wb.save(str(output_file))


import json


def _normalize_docx_spec(raw_spec: Any) -> dict:
    """Normalizes dict/string specs emitted by LLMs into valid DocxSpec schema."""
    if isinstance(raw_spec, str):
        try:
            raw_spec = json.loads(raw_spec)
        except Exception:
            raw_spec = {"title": "Document Deliverable", "sections": [{"body": str(raw_spec)}]}

    if not isinstance(raw_spec, dict):
        raw_spec = {"title": "Document Deliverable", "sections": []}

    title = raw_spec.get("title") or raw_spec.get("document_title") or "Document Deliverable"
    raw_sections = raw_spec.get("sections") or raw_spec.get("blocks") or raw_spec.get("content") or []
    if not isinstance(raw_sections, list):
        raw_sections = [raw_sections]

    normalized_sections = []
    for sec in raw_sections:
        if isinstance(sec, str):
            normalized_sections.append({"body": sec})
        elif isinstance(sec, dict):
            if "headers" in sec or "rows" in sec or "table" in sec:
                headers = sec.get("headers") or []
                rows = sec.get("rows") or []
                normalized_sections.append({
                    "headers": [str(h) for h in headers],
                    "rows": [[str(c) for c in r] if isinstance(r, list) else [str(r)] for r in rows]
                })
            elif "path" in sec or "image_path" in sec:
                img_p = sec.get("path") or sec.get("image_path") or ""
                normalized_sections.append({
                    "path": str(img_p),
                    "caption": str(sec.get("caption")) if sec.get("caption") else None
                })
            else:
                body_val = sec.get("body") or sec.get("content") or sec.get("text") or sec.get("description") or sec.get("details") or ""
                heading_val = sec.get("heading") or sec.get("title") or sec.get("section") or sec.get("header")
                normalized_sections.append({
                    "heading": str(heading_val) if heading_val else None,
                    "body": str(body_val)
                })

    return {"title": str(title), "sections": normalized_sections}


def _normalize_pptx_spec(raw_spec: Any) -> dict:
    """Normalizes dict/string/list specs emitted by LLMs into valid PptxSpec schema."""
    if isinstance(raw_spec, str):
        try:
            raw_spec = json.loads(raw_spec)
        except Exception:
            raw_spec = {"slides": [{"title": "Presentation Deliverable", "blocks": [{"body": str(raw_spec)}]}]}

    if isinstance(raw_spec, list):
        raw_spec = {"slides": raw_spec}

    if not isinstance(raw_spec, dict):
        raw_spec = {"slides": []}

    slides_list = raw_spec.get("slides") or raw_spec.get("sections") or []
    if not isinstance(slides_list, list):
        slides_list = [slides_list]

    normalized_slides = []
    for s_idx, slide_data in enumerate(slides_list):
        if isinstance(slide_data, str):
            normalized_slides.append({
                "title": f"Slide {s_idx + 1}",
                "blocks": [{"body": slide_data}]
            })
        elif isinstance(slide_data, dict):
            stitle = slide_data.get("title") or slide_data.get("heading") or f"Slide {s_idx + 1}"
            blocks_raw = slide_data.get("blocks") or slide_data.get("content") or slide_data.get("sections") or []
            if not isinstance(blocks_raw, list):
                blocks_raw = [blocks_raw]

            normalized_blocks = []
            for b in blocks_raw:
                if isinstance(b, str):
                    normalized_blocks.append({"body": b})
                elif isinstance(b, dict):
                    if "headers" in b or "rows" in b:
                        normalized_blocks.append({
                            "headers": [str(h) for h in b.get("headers", [])],
                            "rows": [[str(c) for c in r] if isinstance(r, list) else [str(r)] for r in b.get("rows", [])]
                        })
                    elif "path" in b or "image_path" in b:
                        normalized_blocks.append({
                            "path": str(b.get("path") or b.get("image_path") or ""),
                            "caption": str(b.get("caption")) if b.get("caption") else None
                        })
                    else:
                        b_head = b.get("heading") or b.get("title")
                        b_body = b.get("body") or b.get("content") or b.get("text") or ""
                        normalized_blocks.append({
                            "heading": str(b_head) if b_head else None,
                            "body": str(b_body)
                        })
            normalized_slides.append({
                "title": str(stitle),
                "blocks": normalized_blocks
            })

    return {"slides": normalized_slides}


@audited_tool
def doc_gen(input: DocGenInput) -> DocGenResult:
    fmt = input.format.lower().strip()
    if fmt not in {"docx", "pptx", "xlsx"}:
        return DocGenResult(
            status=ToolStatus.ERROR,
            error=f"Unsupported format '{input.format}'. Allowed: docx, pptx, xlsx"
        )

    # Validate output path in workspace boundary
    try:
        output_file = validate_workspace_path(input.output_path, input.session_id, create_parents=True)
    except Exception as err:
        return DocGenResult(status=ToolStatus.ERROR, error=str(err))

    try:
        if fmt == "docx":
            norm_spec = _normalize_docx_spec(input.spec)
            spec_obj = DocxSpec.model_validate(norm_spec)
            _render_docx(spec_obj, output_file, input.session_id)

        elif fmt == "pptx":
            norm_spec = _normalize_pptx_spec(input.spec)
            spec_obj = PptxSpec.model_validate(norm_spec)
            _render_pptx(spec_obj, output_file, input.session_id)

        elif fmt == "xlsx":
            spec_obj = XlsxSpec.model_validate(input.spec)
            _render_xlsx(spec_obj, output_file)

        return DocGenResult(
            status=ToolStatus.SUCCESS,
            output_path=str(output_file),
            format=fmt
        )

    except Exception as e:
        return DocGenResult(
            status=ToolStatus.ERROR,
            error=f"Failed generating '{fmt}' document: {str(e)}"
        )
